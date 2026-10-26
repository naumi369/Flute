import streamlit as st
import tempfile
import zipfile
import os

import openpyxl
from generate_order_files import process_file, merge_order_files

st.set_page_config(
    page_title="Order File Generator",
    page_icon="📋",
    layout="centered",
)

st.title("Order File Generator")
st.caption("Generate order files · Paste data · Merge multiple orders for print")

REQUIRED_HEADERS = [
    "Order No", "Style#", "Units", "Stones#", "Size", "Qlty", "Sub Bag#",
]
HEADER_ALIASES = {
    "order no": "Order No", "orderno": "Order No",
    "style#": "Style#", "style #": "Style#", "style": "Style#",
    "units": "Units", "stones#": "Stones#", "stones": "Stones#",
    "size": "Size", "qlty": "Qlty", "quality": "Qlty",
    "sub bag#": "Sub Bag#", "sub bag": "Sub Bag#", "subbag#": "Sub Bag#", "subbag": "Sub Bag#",
    "ship to": "Ship To", "due_date": "Due_date", "due date": "Due_date",
}


def normalize_header(h: str) -> str:
    key = (h or "").strip().lower()
    return HEADER_ALIASES.get(key, (h or "").strip())


def paste_to_xlsx(text: str, dest_path: str) -> tuple:
    text = text.strip()
    if not text:
        raise ValueError("Nothing pasted.")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 2:
        raise ValueError("Need a header row and at least one data row.")
    sample = lines[0]
    delim = "\t" if sample.count("\t") >= sample.count(",") else ","
    rows = []
    for ln in lines:
        parts = [p.strip().strip('"') for p in ln.split(delim)]
        rows.append(parts)
    max_cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < max_cols:
            r.append("")
    headers = [normalize_header(h) for h in rows[0]]
    rows[0] = headers
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValueError(
            "Missing required column(s): " + ", ".join(missing)
            + "\n\nFound headers: " + ", ".join(headers)
        )
    wb = openpyxl.Workbook()
    ws = wb.active
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            if r_idx > 1 and val != "":
                try:
                    if "." in str(val) or (isinstance(val, str) and "e" in val.lower()):
                        val = float(val)
                    else:
                        val = int(val)
                except ValueError:
                    pass
            ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(dest_path)
    return len(rows) - 1, headers


def preview_orders_from_xlsx(path: str) -> list:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_map = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        header_map[i] = normalize_header(str(h))
    from collections import defaultdict
    orders = defaultdict(lambda: {"rows": 0, "style": "", "units": ""})
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        rec = {}
        for i, val in enumerate(row):
            if i in header_map:
                rec[header_map[i]] = val
        ono = rec.get("Order No")
        if not ono:
            continue
        ono = str(ono).strip()
        orders[ono]["rows"] += 1
        if not orders[ono]["style"]:
            orders[ono]["style"] = rec.get("Style#") or ""
        if not orders[ono]["units"]:
            orders[ono]["units"] = rec.get("Units") or ""
    return [
        {"Order No": k, "Style#": v["style"], "Units": v["units"], "Lines": v["rows"]}
        for k, v in sorted(orders.items())
    ]


def run_generation(input_path: str):
    with tempfile.TemporaryDirectory() as tmp:
        out_dir = os.path.join(tmp, "output")
        os.makedirs(out_dir, exist_ok=True)
        logs = []
        def log_cb(msg):
            logs.append(msg)
        created = process_file(input_path, out_dir, log_callback=log_cb)
        st.text_area("Log", value="\n".join(logs), height=220)
        if created:
            zip_path = os.path.join(tmp, "orders.zip")
            with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
                for p in created:
                    zf.write(p, arcname=os.path.basename(p))
            with open(zip_path, "rb") as f:
                st.download_button(
                    label=f"⬇ Download {len(created)} order file(s) as ZIP",
                    data=f.read(),
                    file_name="order_files.zip",
                    mime="application/zip",
                )
            st.success(f"Created {len(created)} file(s).")
        else:
            st.error("No files were created. Check the log above.")


tab_upload, tab_paste, tab_merge = st.tabs([
    "📁 Upload Excel file",
    "📋 Paste from Excel",
    "📄 Merge order files",
])

# ----- Upload -----
with tab_upload:
    uploaded = st.file_uploader(
        "Daily sample data (.xlsx)",
        type=["xlsx", "xls"],
        key="uploader",
    )
    if uploaded is not None:
        st.success(f"Loaded: **{uploaded.name}** ({uploaded.size:,} bytes)")
        if st.button("Generate order files", type="primary", key="gen_upload"):
            with st.spinner("Generating…"):
                with tempfile.TemporaryDirectory() as tmp:
                    input_path = os.path.join(tmp, uploaded.name)
                    with open(input_path, "wb") as f:
                        f.write(uploaded.getbuffer())
                    try:
                        summary = preview_orders_from_xlsx(input_path)
                        if summary:
                            st.markdown("**Orders detected**")
                            st.dataframe(summary, use_container_width=True, hide_index=True)
                    except Exception:
                        pass
                    run_generation(input_path)
    else:
        st.info("Upload your daily sample Excel file to begin.")

# ----- Paste -----
with tab_paste:
    st.markdown(
        """
        **How to paste**
        1. In Excel, select data **including the header row**
        2. Copy (`Ctrl+C` / `Cmd+C`)
        3. Paste below (`Ctrl+V` / `Cmd+V`)
        """
    )
    st.caption("Required columns: " + ", ".join(f"`{h}`" for h in REQUIRED_HEADERS))
    pasted = st.text_area(
        "Paste Excel data here",
        height=260,
        placeholder="Ship To\tStyle#\tOrder No\t...\n100050\t31-D36867SY-E.00\tP2009515\t...",
        key="paste_area",
    )
    if pasted and pasted.strip():
        n_lines = len([ln for ln in pasted.strip().splitlines() if ln.strip()])
        st.caption(f"Detected **{max(0, n_lines - 1)}** data row(s) (+ 1 header).")
        if st.button("Generate order files", type="primary", key="gen_paste"):
            with st.spinner("Validating paste and generating…"):
                try:
                    with tempfile.TemporaryDirectory() as tmp:
                        input_path = os.path.join(tmp, "pasted_data.xlsx")
                        n, headers = paste_to_xlsx(pasted, input_path)
                        st.success(
                            f"Validated **{n}** data row(s). Headers OK: "
                            + ", ".join(headers[:8])
                            + ("…" if len(headers) > 8 else "")
                        )
                        summary = preview_orders_from_xlsx(input_path)
                        if summary:
                            st.markdown("**Orders detected**")
                            st.dataframe(summary, use_container_width=True, hide_index=True)
                        run_generation(input_path)
                except ValueError as e:
                    st.error(str(e))
                except Exception as e:
                    st.error(f"Could not process pasted data: {e}")
    else:
        st.info("Copy rows from Excel (with headers) and paste them above.")

# ----- Merge -----
with tab_merge:
    st.markdown(
        """
        **Merge multiple order Excel files** into one printable workbook.

        - Files are stacked **one after another** with **2 blank rows** between
        - **No forced page break** between files — if an order fits in half a page,
          the next order continues on the **same page** (no blank half-page)
        - Print: **A4 · Portrait · Narrow margins · Fit to page width**
        """
    )
    merge_files = st.file_uploader(
        "Select multiple order Excel files",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        key="merge_uploader",
        help="Upload the generated order files (e.g. P2009515.xlsx, P2009516.xlsx, …)",
    )
    if merge_files:
        st.write(f"**{len(merge_files)}** file(s) selected:")
        for f in merge_files:
            st.caption(f"• {f.name} ({f.size:,} bytes)")

        # Optional: let user set order by sorting names
        sort_alpha = st.checkbox("Sort files alphabetically by name", value=True)

        if st.button("Merge for print", type="primary", key="do_merge"):
            with st.spinner("Merging…"):
                try:
                    files = list(merge_files)
                    if sort_alpha:
                        files = sorted(files, key=lambda x: x.name.lower())

                    with tempfile.TemporaryDirectory() as tmp:
                        paths = []
                        for f in files:
                            p = os.path.join(tmp, f.name)
                            with open(p, "wb") as out:
                                out.write(f.getbuffer())
                            paths.append(p)

                        out_path = os.path.join(tmp, "merged_orders.xlsx")
                        result_path, n_files, n_rows = merge_order_files(paths, out_path)

                        st.success(
                            f"Merged **{n_files}** file(s) into one sheet "
                            f"({n_rows} rows). Short forms will share a page when they fit."
                        )
                        with open(result_path, "rb") as f:
                            st.download_button(
                                label="⬇ Download merged Excel",
                                data=f.read(),
                                file_name="merged_orders.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            )
                except Exception as e:
                    st.error(f"Merge failed: {e}")
    else:
        st.info("Upload two or more order Excel files to merge them for printing.")

st.divider()
st.caption("Output: A4 · Portrait · Narrow margins · Fit to page width · Auto-flow (no wasted half pages)")
